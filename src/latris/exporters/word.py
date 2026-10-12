"""Word (.docx) report exporter with structured sections and tables.

Uses python-docx. Mirrors the markdown exporter's structure but produces
a Word document with headings, paragraphs, and tables.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from latris.exporters._format import BADGE_TEXT as _BADGE_TEXT
from latris.exporters._format import BADGE_URL as _BADGE_URL
from latris.exporters._format import actor_name as _actor_name
from latris.exporters._scope import view_member_ids
from latris.exporters.packs import get_pack
from latris.repo import git as git_repo
from latris.repo.layout import ProjectPaths
from latris.repo.yaml_io import iter_jsonl, load_model
from latris.schema import (
    FINDING_CATEGORIES,
    AnalysisRun,
    AssumptionLog,
    ChallengeLog,
    ConfidenceSet,
    DataActionLog,
    DatasetLibrary,
    DecisionLog,
    DescriptionSet,
    ElicitationLog,
    ElicitationSession,
    EvidenceLinkSet,
    ExclusionLog,
    MilestoneLog,
    PolaritySet,
    PredictionLog,
    QuestionLog,
    ReferenceLibrary,
    SessionLog,
    SignoffLog,
    Substrate,
)
from latris.schema.ids import id_timestamp


def render(
    paths: ProjectPaths,
    *,
    pack: str | None = None,
    sections: list[str] | None = None,
    model_type: str | None = None,
) -> Document:
    """Build a Word Document from the project, optionally filtered by pack or
    sections. When ``model_type`` is given, the substrate is scoped to that
    formalism's view and findings are filtered to that model type."""
    included = None
    if sections:
        included = set(sections)
    elif pack:
        included = set(get_pack(pack).sections)

    def _include(section: str) -> bool:
        return included is None or section in included

    substrate = load_model(Substrate, paths.conceptual_yaml)
    members = view_member_ids(paths, model_type)
    if members is not None:
        substrate.entities = [e for e in substrate.entities if e.id in members]
        substrate.edges = [e for e in substrate.edges if e.id in members]
    descriptions = {(e.target_id, e.attribute): e for e in load_model(DescriptionSet, paths.descriptions_yaml).entries}
    confidence = {(e.target_id, e.attribute): e for e in load_model(ConfidenceSet, paths.confidence_yaml).entries}
    pol_path = paths.view_attribute_path("cld", "polarity")
    pol_set = load_model(PolaritySet, pol_path) if pol_path.exists() else PolaritySet()
    polarity = {(e.target_id, e.attribute): e for e in pol_set.entries}
    evidence = load_model(EvidenceLinkSet, paths.evidence_links_yaml)
    references = load_model(ReferenceLibrary, paths.references_yaml).entries
    assumptions = load_model(AssumptionLog, paths.assumptions_yaml).entries
    decisions = load_model(DecisionLog, paths.decisions_yaml).entries
    signoffs = load_model(SignoffLog, paths.signoffs_yaml).entries
    sessions = load_model(SessionLog, paths.sessions_yaml).entries
    questions = load_model(QuestionLog, paths.questions_yaml).entries
    exclusions = load_model(ExclusionLog, paths.exclusions_yaml).entries
    predictions = load_model(PredictionLog, paths.predictions_yaml).entries
    challenges = load_model(ChallengeLog, paths.challenges_yaml).entries
    datasets = load_model(DatasetLibrary, paths.data_datasets_yaml).entries
    elicit_index = load_model(ElicitationLog, paths.elicitations_yaml)
    elicit_sessions: list[ElicitationSession] = []
    for entry in elicit_index.entries:
        session_path = paths.elicitations_dir / entry.file_path
        if session_path.exists():
            elicit_sessions.append(load_model(ElicitationSession, session_path))
    from latris.schema.findings import Finding

    analysis_runs = list(iter_jsonl(AnalysisRun, paths.analysis_runs_jsonl))
    all_findings = list(iter_jsonl(Finding, paths.findings_jsonl))
    if model_type is not None:
        all_findings = [f for f in all_findings if f.model_type == model_type]
    milestones = (
        load_model(MilestoneLog, paths.milestones_yaml) if paths.milestones_yaml.exists() else MilestoneLog()
    ).entries

    labels = {n.id: n.label for n in substrate.entities}
    doc = Document()

    # Foundation + Narrative Preamble
    if _include("foundation"):
        from latris.exporters.narrative import build_narrative

        narrative = build_narrative(paths, model_type=model_type)
        doc.add_heading(narrative.title, level=1)
        doc.add_paragraph(narrative.purpose)
        if narrative.boundary:
            doc.add_paragraph(narrative.boundary)

        # Model structure (AI prose or mechanical fallback)
        if narrative.structure_description:
            doc.add_heading("Model Structure", level=2)
            doc.add_paragraph(narrative.structure_description)

        # Model history (AI prose or mechanical fallback)
        if narrative.history_summary:
            doc.add_heading("Model History", level=2)
            doc.add_paragraph(narrative.history_summary)

    if _include("modelling_approach"):
        _render_modelling_approach_docx(doc, paths)

    # Nodes
    if _include("nodes"):
        doc.add_heading("Nodes", level=2)
        if substrate.entities:
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["Label", "Kind", "Description", "Confidence", "Maturity"]):
                hdr[i].text = h
            for n in substrate.entities:
                row = table.add_row().cells
                row[0].text = n.label
                row[1].text = n.kind
                desc = descriptions.get((n.id, "description"))
                row[2].text = desc.value if desc else ""
                conf = confidence.get((n.id, "confidence"))
                row[3].text = conf.value if conf else ""
                row[4].text = n.maturity
        else:
            doc.add_paragraph("No nodes yet.")

    # Edges
    if _include("edges"):
        doc.add_heading("Edges", level=2)
        if substrate.edges:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["From → To", "Kind", "Polarity", "Confidence"]):
                hdr[i].text = h
            for e in substrate.edges:
                row = table.add_row().cells
                src = labels.get(e.from_id, e.from_id)
                dst = labels.get(e.to_id, e.to_id)
                row[0].text = f"{src} → {dst}"
                row[1].text = e.kind
                pol = polarity.get((e.id, "polarity"))
                row[2].text = pol.value if pol else ""
                conf = confidence.get((e.id, "confidence"))
                row[3].text = conf.value if conf else ""
        else:
            doc.add_paragraph("No edges yet.")

    # Evidence
    if _include("evidence"):
        doc.add_heading("Evidence", level=2)
        if evidence.entries:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["Target", "Attribute", "Reference", "Notes"]):
                hdr[i].text = h
            for link in evidence.entries:
                row = table.add_row().cells
                row[0].text = labels.get(link.target_id, link.target_id)
                row[1].text = link.attribute
                row[2].text = link.bibkey
                row[3].text = link.notes or ""
        else:
            doc.add_paragraph("No evidence attached yet.")

    # Assumptions
    if _include("assumptions"):
        doc.add_heading("Assumptions", level=2)
        if assumptions:
            for a in assumptions:
                doc.add_paragraph(
                    f"{a.statement} — by {_actor_name(a.authored_by)}, {id_timestamp(a.id).date().isoformat()}",
                    style="List Bullet",
                )
        else:
            doc.add_paragraph("None.")

    # Decisions
    if _include("decisions"):
        doc.add_heading("Decisions", level=2)
        if decisions:
            for d in decisions:
                doc.add_paragraph(
                    f"{d.statement} — {d.rationale} — by {_actor_name(d.authored_by)}, "
                    f"{id_timestamp(d.id).date().isoformat()}",
                    style="List Bullet",
                )
        else:
            doc.add_paragraph("None.")

    # Sign-offs
    if _include("signoffs"):
        doc.add_heading("Sign-offs", level=2)
        if signoffs:
            for s in signoffs:
                doc.add_paragraph(
                    f"{s.role} on {s.scope} — {_actor_name(s.reviewer)}, {s.signed_at.date().isoformat()}",
                    style="List Bullet",
                )
        else:
            doc.add_paragraph("None.")

    # Sessions
    if _include("sessions"):
        doc.add_heading("Sessions", level=2)
        if sessions:
            for s in sessions:
                attendees = ", ".join(_actor_name(a) for a in s.attendees)
                doc.add_paragraph(
                    f"{s.held_at.isoformat()} with {attendees}",
                    style="List Bullet",
                )
        else:
            doc.add_paragraph("None.")

    # Questions
    if _include("questions"):
        doc.add_heading("Open Questions", level=2)
        open_qs = [q for q in questions if q.disposition == "open"]
        if open_qs:
            for q in open_qs:
                doc.add_paragraph(
                    f"{q.question} (raised by {_actor_name(q.raised_by)}, {id_timestamp(q.id).date().isoformat()})",
                    style="List Bullet",
                )
        else:
            doc.add_paragraph("None.")

    # Exclusions
    if _include("exclusions"):
        doc.add_heading("Exclusions", level=2)
        if exclusions:
            for e in exclusions:
                doc.add_paragraph(
                    f"{e.statement} — {e.rationale} — by {_actor_name(e.authored_by)}, "
                    f"{id_timestamp(e.id).date().isoformat()}",
                    style="List Bullet",
                )
        else:
            doc.add_paragraph("None.")

    # Predictions
    if _include("predictions"):
        doc.add_heading("Predictions", level=2)
        if predictions:
            for p in predictions:
                doc.add_paragraph(
                    f"{p.statement} (horizon: {p.horizon}; status: {p.resolution}) "
                    f"— by {_actor_name(p.authored_by)}, "
                    f"{id_timestamp(p.id).date().isoformat()}",
                    style="List Bullet",
                )
        else:
            doc.add_paragraph("None.")

    # Challenges
    if _include("challenges"):
        doc.add_heading("Challenges", level=2)
        if challenges:
            for c in challenges:
                target_label = labels.get(c.target_id, c.target_id)
                doc.add_paragraph(
                    f"[{c.state}] {target_label}: {c.argument} "
                    f"(opened by {_actor_name(c.opened_by)}, "
                    f"{id_timestamp(c.id).date().isoformat()})",
                    style="List Bullet",
                )
        else:
            doc.add_paragraph("None raised.")

    # References
    if _include("references"):
        doc.add_heading("References", level=2)
        if references:
            for r in sorted(references, key=lambda r: r.bibkey):
                authors = ", ".join(r.authors) if r.authors else ""
                year = f"({r.year})" if r.year else ""
                doc.add_paragraph(f"[{r.bibkey}] {authors} {year} {r.title}", style="List Bullet")
        else:
            doc.add_paragraph("None.")

    # Elicitation Sessions
    if _include("elicitations"):
        doc.add_heading("Elicitation Sessions", level=2)
        if elicit_sessions:
            for s in elicit_sessions:
                experts = ", ".join(e.name for e in s.experts) if s.experts else "none listed"
                targets = (
                    ", ".join(f"{labels.get(t.target_id, t.target_id)}/{t.attribute}" for t in s.target_attributes)
                    if s.target_attributes
                    else "none"
                )
                notes = f" — {s.notes}" if s.notes else ""
                doc.add_paragraph(
                    f"{s.held_at.isoformat()} — {s.protocol.name} "
                    f"(facilitated by {s.facilitator.name}; "
                    f"experts: {experts}; targets: {targets}; "
                    f"phase: {s.phase}){notes}",
                    style="List Bullet",
                )
        else:
            doc.add_paragraph("None.")

    # Datasets
    if _include("datasets"):
        doc.add_heading("Datasets", level=2)
        if datasets:
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["Name", "Source", "Rows", "Columns", "Description"]):
                hdr[i].text = h
            for d in datasets:
                row = table.add_row().cells
                row[0].text = d.name
                row[1].text = d.source_type
                row[2].text = str(d.rows) if d.rows is not None else "?"
                row[3].text = str(len(d.columns)) if d.columns else "?"
                row[4].text = d.description
        else:
            doc.add_paragraph("None registered.")

    # Data Actions
    if _include("data_actions"):
        doc.add_heading("Data Actions", level=2)
        all_actions = []
        for d in datasets:
            log_path = paths.data_action_log(d.id)
            if log_path.exists():
                log = load_model(DataActionLog, log_path)
                for a in log.entries:
                    all_actions.append((d.name, a))
        if all_actions:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["Dataset", "Action", "Summary", "Date"]):
                hdr[i].text = h
            for dname, a in all_actions:
                row = table.add_row().cells
                row[0].text = dname
                row[1].text = a.kind
                row[2].text = a.summary
                row[3].text = id_timestamp(a.id).date().isoformat()
        else:
            doc.add_paragraph("No data actions recorded.")

    # Analysis Runs
    if _include("analyses"):
        doc.add_heading("Analysis Runs", level=2)
        if analysis_runs:
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["Purpose", "Algorithm", "Run by", "Date", "Summary"]):
                hdr[i].text = h
            for r in analysis_runs:
                row = table.add_row().cells
                row[0].text = r.purpose
                row[1].text = r.algorithm
                row[2].text = _actor_name(r.run_by)
                row[3].text = id_timestamp(r.id).date().isoformat()
                row[4].text = r.output_summary
        else:
            doc.add_paragraph("None.")

    # Findings
    if _include("findings"):
        visible_maturities = {"proposed", "reviewed", "approved"}
        if pack == "client" or pack == "publication":
            visible_maturities = {"approved"}

        visible = [f for f in all_findings if f.maturity in visible_maturities]
        doc.add_heading("Findings", level=2)
        if visible:
            for cat in FINDING_CATEGORIES:
                cat_findings = [f for f in visible if f.category == cat]
                if not cat_findings:
                    continue
                doc.add_heading(cat.title(), level=3)
                for f in cat_findings:
                    draft_marker = " (proposed)" if f.maturity == "proposed" else ""
                    date_str = id_timestamp(f.id).date().isoformat()
                    doc.add_paragraph(
                        f"{f.statement}{draft_marker} — {f.model_type}, {f.maturity}, {date_str}",
                        style="List Bullet",
                    )
                    if f.detail:
                        doc.add_paragraph(f.detail)
        else:
            doc.add_paragraph("No findings recorded.")

    # Interpretation
    if _include("interpretation"):
        _render_interpretation_docx(doc, paths, model_type=model_type)

    # Milestones
    if _include("milestones"):
        doc.add_heading("Milestones", level=2)
        if milestones:
            milestone_dates = git_repo.commit_timestamp_many(
                [m.commit_sha for m in milestones], paths.root
            )
            for m in milestones:
                notes = f" — {m.notes}" if m.notes else ""
                date = milestone_dates.get(m.commit_sha)
                date_str = date.date().isoformat() if date else "—"
                doc.add_paragraph(
                    f"{m.name} ({m.maturity}) — by {_actor_name(m.created_by)}, "
                    f"{date_str}{notes}",
                    style="List Bullet",
                )
        else:
            doc.add_paragraph("None.")

    # Footer badge
    badge_run = doc.add_paragraph().add_run(f"{_BADGE_TEXT} — {_BADGE_URL}")
    badge_run.italic = True

    return doc


def write(
    paths: ProjectPaths,
    target: Path,
    *,
    pack: str | None = None,
    model_type: str | None = None,
) -> Path:
    """Render and save a Word document."""
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = render(paths, pack=pack, model_type=model_type)
    doc.save(str(target))
    return target


def _render_interpretation_docx(doc, paths: ProjectPaths, *, model_type: str | None = None) -> None:
    """Add the Interpret-phase sections (Answer, Confidence, Limitations, etc).

    No-op if interpretation.yaml is missing or has no sections — keeps the
    report quiet for projects that haven't reached the Interpret phase. Mirrors
    the Markdown exporter's ``_render_interpretation`` (registry ordering, model
    suffix disambiguation)."""
    interp_path = paths.interpretation_yaml
    if not interp_path.exists():
        return
    from latris.operations.interpret_registry import (
        INTERPRET_SECTION_REGISTRY,
        get_section_spec,
    )
    from latris.schema.interpretation import InterpretationStore

    store = load_model(InterpretationStore, interp_path)
    scoped_sections = [s for s in store.sections if model_type is None or s.model_type == model_type]
    if not scoped_sections:
        return

    doc.add_heading("Interpretation", level=2)

    # Render in registry order so the report flows logically
    # (Structural Summary → Key Findings → ... → Open Questions).
    by_type: dict[str, list] = {}
    for s in scoped_sections:
        by_type.setdefault(s.section_type, []).append(s)

    for section_type in INTERPRET_SECTION_REGISTRY:
        for s in by_type.get(section_type, []):
            try:
                title = get_section_spec(s.section_type).display_name
            except ValueError:
                title = s.section_type.replace("_", " ").title()

            # Disambiguate by model only in the aggregate (unscoped) report.
            multi = len(by_type.get(section_type, [])) > 1 or len(scoped_sections) > 10
            heading_suffix = f" ({s.model_type.upper()})" if model_type is None and multi else ""
            doc.add_heading(f"{title}{heading_suffix}", level=3)

            if s.prose:
                doc.add_paragraph(s.prose.rstrip())
            elif s.claims:
                doc.add_paragraph("AI prose not generated yet — claims only:")
            else:
                doc.add_paragraph("Not yet populated.")

            if s.claims:
                doc.add_paragraph("Claims (from analysis):")
                for c in s.claims:
                    doc.add_paragraph(c, style="List Bullet")

            state_marker = {
                "empty": "empty",
                "claims_only": "claims only",
                "generated": "AI-generated",
                "edited": "human-edited",
            }.get(s.generation_state, s.generation_state)
            doc.add_paragraph(f"State: {state_marker} · last updated {s.updated_at.date().isoformat()}")


def _render_modelling_approach_docx(doc, paths: ProjectPaths) -> None:
    """Add the 'Modelling Approach' section (why each formalism was chosen)
    from the ModelChoiceLog. No-op if no choices are recorded."""
    from latris.operations.project import current_model_choices
    from latris.schema.model_types import get_model_type

    choices = current_model_choices(root=paths.root)
    if not choices:
        return
    doc.add_heading("Modelling Approach", level=2)
    if len(choices) > 1:
        doc.add_paragraph(
            "This question was modelled with more than one formalism to triangulate "
            "the answer; each choice is justified below."
        )
    for mt_id, choice in choices.items():
        try:
            name = get_model_type(mt_id).display_name
        except ValueError:
            name = mt_id.upper()
        doc.add_heading(name, level=3)
        doc.add_paragraph(choice.rationale.strip() or "No rationale recorded for this choice.")

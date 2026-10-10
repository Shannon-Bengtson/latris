"""Diagram embedding helpers — render the model graph for export reports.

LIVE: ``dot_source`` and ``render_svg`` back the HTML exporters (review_html,
story_html). ``render_svg`` is graphviz's plain SVG; story mode prefers its own
``render_story_svg`` (exporters/story_html.py) and only falls back to this.

LEGACY / not currently wired: ``render_markdown_diagram`` (fenced DOT block),
``render_png_bytes``, and ``add_diagram_to_word`` (graphviz PNG for Word). The
Markdown and Word exporters ship no diagram today. The intended long-term
direction is to converge report diagram embedding on the story-mode SVG
renderer (``render_story_svg``) rather than these graphviz PNG/DOT paths — kept
here, tested, until that unification lands. Tracked as a separate feature.
"""

from __future__ import annotations

import subprocess
from io import BytesIO

from latris.exporters.dot import render as dot_render
from latris.repo.layout import ProjectPaths


def dot_source(paths: ProjectPaths, view: str = "system_map") -> str:
    """Return DOT source for a view (whole substrate by default)."""
    return dot_render(paths, view=view)  # type: ignore[arg-type]


def render_markdown_diagram(paths: ProjectPaths) -> str:
    """Return a Markdown-embeddable diagram section."""
    source = dot_source(paths)
    if "// empty project" in source:
        return ""
    return f"## Model Diagram\n\n```dot\n{source}```\n"


def render_svg(paths: ProjectPaths, view: str = "system_map") -> str | None:
    """Try to render DOT to SVG via graphviz. Returns None if unavailable."""
    source = dot_source(paths, view=view)
    if "// empty project" in source:
        return None
    try:
        result = subprocess.run(
            ["dot", "-Tsvg"],
            input=source.encode("utf-8"),
            capture_output=True,
            timeout=10,
        )
        if result.returncode == 0 and result.stdout:
            return result.stdout.decode("utf-8")
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass
    return None


def render_png_bytes(paths: ProjectPaths) -> bytes | None:
    """Try to render DOT to PNG via graphviz. Returns None if unavailable."""
    source = dot_source(paths)
    if "// empty project" in source:
        return None
    try:
        result = subprocess.run(
            ["dot", "-Tpng"],
            input=source.encode("utf-8"),
            capture_output=True,
            timeout=10,
        )
        if result.returncode == 0 and result.stdout:
            return result.stdout
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass
    return None


def add_diagram_to_word(doc, paths: ProjectPaths) -> None:
    """Add a model diagram to a Word document.

    Tries graphviz PNG first; falls back to a text-based summary.
    """
    from docx.shared import Inches

    png = render_png_bytes(paths)
    if png is not None:
        doc.add_heading("Model Diagram", level=2)
        doc.add_picture(BytesIO(png), width=Inches(6.0))
        return

    # Fallback: text-based graph summary
    from latris.repo.yaml_io import load_model
    from latris.schema import Substrate

    substrate = load_model(Substrate, paths.conceptual_yaml)
    if not substrate.entities and not substrate.edges:
        return

    doc.add_heading("Model Diagram", level=2)
    doc.add_paragraph(
        "Install Graphviz (graphviz.org) to render a visual diagram. The model structure is summarised below.",
        style="Body Text",
    )

    labels = {e.id: e.label for e in substrate.entities}
    for edge in substrate.edges:
        src = labels.get(edge.from_id, edge.from_id[:12])
        dst = labels.get(edge.to_id, edge.to_id[:12])
        doc.add_paragraph(f"{src}  \u2192  {dst}  ({edge.kind})", style="List Bullet")

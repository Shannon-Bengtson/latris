"""Tests for Word (.docx) export (M15)."""

from __future__ import annotations

from pathlib import Path

import pytest

from latris.exporters import word as word_exporter
from latris.operations import edges as edge_ops
from latris.operations import nodes as node_ops
from latris.operations.init import init_project
from latris.repo.layout import ProjectPaths


@pytest.fixture
def project(tmp_path: Path, git_identity, monkeypatch):
    root = tmp_path / "demo"
    init_project(root, statement="Does X cause Y?")
    a, _ = node_ops.add_node("Price", kind="variable", root=root)
    b, _ = node_ops.add_node("Demand", kind="variable", root=root)
    edge_ops.add_edge(a.id, b.id, kind="causal", root=root)
    monkeypatch.chdir(root)
    return root


def test_word_render_produces_document(project):
    paths = ProjectPaths(project)
    doc = word_exporter.render(paths)
    # Document has paragraphs
    assert hasattr(doc, "paragraphs")
    assert len(doc.paragraphs) > 0


def test_word_write_creates_valid_docx(project, tmp_path):
    paths = ProjectPaths(project)
    output = tmp_path / "output.docx"
    result = word_exporter.write(paths, output)
    assert result.exists()
    assert result.suffix == ".docx"
    # Verify it can be loaded back by python-docx
    from docx import Document as load_docx

    loaded = load_docx(str(result))
    assert len(loaded.paragraphs) > 0


def test_word_with_pack_filters_sections(project, tmp_path):
    paths = ProjectPaths(project)

    # Full export
    full_doc = word_exporter.render(paths)
    full_headings = [p.text for p in full_doc.paragraphs if p.style.name.startswith("Heading")]

    # Publication pack — no challenges, predictions, sessions, etc.
    pub_doc = word_exporter.render(paths, pack="publication")
    pub_headings = [p.text for p in pub_doc.paragraphs if p.style.name.startswith("Heading")]

    assert "Challenges" in full_headings
    assert "Challenges" not in pub_headings
    assert "Nodes" in pub_headings
    assert "References" in pub_headings


def test_word_contains_node_data(project):
    paths = ProjectPaths(project)
    doc = word_exporter.render(paths)
    # Find all text in the document
    all_text = "\n".join(p.text for p in doc.paragraphs)
    # Tables aren't in paragraphs, check them separately
    all_table_text = ""
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_table_text += cell.text + " "
    combined = all_text + all_table_text
    assert "Price" in combined
    assert "Demand" in combined


def test_word_renders_interpretation_section(project):
    """Interpret-phase prose flows into the Word report (parity with Markdown)."""
    from latris.repo.yaml_io import dump_model
    from latris.schema.ids import new_id
    from latris.schema.interpretation import InterpretationStore, InterpretSection

    paths = ProjectPaths(project)
    store = InterpretationStore(
        sections=[
            InterpretSection(
                id=new_id("interpret-section"),
                section_type="structural_summary",
                model_type="cld",
                prose="The system shows a reinforcing loop.",
                generation_state="edited",
            )
        ]
    )
    dump_model(store, paths.interpretation_yaml)

    doc = word_exporter.render(paths)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    body = "\n".join(p.text for p in doc.paragraphs)

    assert "Interpretation" in headings
    assert "The system shows a reinforcing loop." in body


def test_word_footer_badge(project):
    """Document ends with an italic 'Made with Latris' line (parity with the
    Markdown exporter's rule + italic footer)."""
    paths = ProjectPaths(project)
    doc = word_exporter.render(paths)
    last_para = doc.paragraphs[-1]
    assert last_para.text == "Made with Latris — https://latris.app/?utm_source=latris-export&utm_medium=badge"
    assert last_para.runs[0].italic is True
